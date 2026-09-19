"""Importacion del checklist desde Excel.

Los libros de prueba se generan al vuelo con openpyxl para no depender de un
archivo binario en el repositorio.
"""

from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import Workbook

from mukuwareru.nucleo.repositorios import RepositorioProyectos
from mukuwareru.nucleo.servicios import ServicioImportacion, ServicioProgreso

EXCEL_REAL = Path(__file__).resolve().parents[1] / "CFA.xlsx"


@pytest.fixture
def proyecto(conn: sqlite3.Connection) -> int:
    return RepositorioProyectos(conn).crear("CFA Level I").id


def _libro(tmp_path: Path) -> Path:
    """Replica la forma del Excel real: cabeceras desplazadas y fila de totales."""
    wb = Workbook()

    hoja = wb.active
    hoja.title = "Módulos"
    hoja["B2"] = "Checklist de Módulos"
    hoja["B5"], hoja["C5"], hoja["D5"] = "Tema", "Módulo / Reading", "Completado"
    filas = [
        ("Ethics", "Code of Ethics", True),
        ("Ethics", "Guidance", False),
        ("Quantitative Methods", "Rates and Returns", True),
        ("Total completados / total módulos", None, None),
    ]
    for indice, (tema, modulo, hecho) in enumerate(filas, start=6):
        hoja[f"B{indice}"], hoja[f"C{indice}"], hoja[f"D{indice}"] = tema, modulo, hecho

    registro = wb.create_sheet("Registro Diario")
    registro["B4"], registro["C4"] = "Fecha", "Horas"
    registro["B5"], registro["C5"] = datetime(2026, 8, 11), 2.0
    registro["B6"], registro["C6"] = datetime(2026, 8, 12), 2.5
    registro["B7"], registro["C7"] = "TOTAL HORAS", 4.5

    ruta = tmp_path / "progreso.xlsx"
    wb.save(ruta)
    return ruta


# --- Analisis --------------------------------------------------------------


def test_analizar_encuentra_las_hojas_por_cabecera(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    informe = ServicioImportacion(conn).analizar(_libro(tmp_path))
    assert informe.hojas_ausentes == []
    assert len(informe.modulos) == 3
    assert informe.completados == 2
    assert informe.materias == ["Ethics", "Quantitative Methods"]


def test_analizar_descarta_las_filas_de_totales(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    informe = ServicioImportacion(conn).analizar(_libro(tmp_path))
    assert len(informe.descartadas) == 2
    assert all("otales" in d or "TOTAL" in d for d in informe.descartadas)


def test_analizar_convierte_horas_en_segundos(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    informe = ServicioImportacion(conn).analizar(_libro(tmp_path))
    assert [s.segundos for s in informe.sesiones] == [7200, 9000]
    assert informe.horas == 4.5


def test_analizar_no_escribe_nada(
    conn: sqlite3.Connection, tmp_path: Path, proyecto: int
) -> None:
    ServicioImportacion(conn).analizar(_libro(tmp_path))
    assert conn.execute("SELECT COUNT(*) FROM modulo").fetchone()[0] == 0
    assert conn.execute("SELECT COUNT(*) FROM sesion").fetchone()[0] == 0


# --- Aplicacion ------------------------------------------------------------


def test_aplicar_crea_materias_modulos_y_sesiones(
    conn: sqlite3.Connection, tmp_path: Path, proyecto: int
) -> None:
    servicio = ServicioImportacion(conn)
    resultado = servicio.aplicar(proyecto, servicio.analizar(_libro(tmp_path)))

    assert resultado.materias_creadas == 2
    assert resultado.modulos_creados == 3
    assert resultado.sesiones_creadas == 2

    avance = ServicioProgreso(conn).resumen(proyecto)
    assert (avance.total, avance.completados) == (3, 2)


def test_reimportar_no_duplica_nada(
    conn: sqlite3.Connection, tmp_path: Path, proyecto: int
) -> None:
    servicio = ServicioImportacion(conn)
    informe = servicio.analizar(_libro(tmp_path))
    servicio.aplicar(proyecto, informe)
    segunda = servicio.aplicar(proyecto, informe)

    assert segunda.materias_creadas == 0
    assert segunda.modulos_creados == 0
    assert segunda.sesiones_creadas == 0
    assert segunda.sesiones_omitidas == 2
    assert conn.execute("SELECT COUNT(*) FROM sesion").fetchone()[0] == 2
    assert conn.execute("SELECT COUNT(*) FROM modulo").fetchone()[0] == 3


def test_reimportar_actualiza_un_modulo_marcado_despues(
    conn: sqlite3.Connection, tmp_path: Path, proyecto: int
) -> None:
    servicio = ServicioImportacion(conn)
    ruta = _libro(tmp_path)
    servicio.aplicar(proyecto, servicio.analizar(ruta))

    informe = servicio.analizar(ruta)
    next(m for m in informe.modulos if m.nombre == "Guidance").completado = True
    resultado = servicio.aplicar(proyecto, informe)

    assert resultado.modulos_actualizados == 1
    assert ServicioProgreso(conn).resumen(proyecto).completados == 3


def test_una_hoja_ausente_se_reporta_en_vez_de_fallar(
    conn: sqlite3.Connection, tmp_path: Path
) -> None:
    wb = Workbook()
    wb.active.title = "Otra cosa"
    wb.active["A1"] = "nada util"
    ruta = tmp_path / "vacio.xlsx"
    wb.save(ruta)

    informe = ServicioImportacion(conn).analizar(ruta)
    assert set(informe.hojas_ausentes) == {"modulos", "registro diario"}
    assert informe.modulos == []


# --- Contra el archivo real del usuario ------------------------------------


@pytest.mark.skipif(not EXCEL_REAL.exists(), reason="CFA.xlsx no esta en el repositorio")
def test_el_excel_real_se_importa_completo(
    conn: sqlite3.Connection, proyecto: int
) -> None:
    servicio = ServicioImportacion(conn)
    informe = servicio.analizar(EXCEL_REAL)

    # 93 modulos en 10 temas. El numero sigue al curriculo del CFA, que se
    # renumera cada ano: si cambia el temario, se actualiza aqui a proposito.
    # Lo que de verdad vigila esta prueba es que las filas de totales de la hoja
    # no se cuelen como si fueran modulos.
    assert len(informe.modulos) == 93
    assert len(informe.materias) == 10
    assert not any("total" in m.nombre.lower() for m in informe.modulos)
    assert any(d.startswith("Fila de totales") for d in informe.descartadas)

    servicio.aplicar(proyecto, informe)
    avance = ServicioProgreso(conn).resumen(proyecto)
    assert avance.total == 93
    assert avance.completados == informe.completados
    assert {m.nombre for m in avance.materias} == set(informe.materias)
