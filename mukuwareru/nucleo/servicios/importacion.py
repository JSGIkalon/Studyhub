"""Importacion del checklist de estudio desde Excel.

Migracion de una sola vez: despues la aplicacion es la fuente oficial. El
servicio se divide en dos pasos deliberadamente, ``analizar`` y ``aplicar``,
para que la interfaz pueda mostrar una vista previa antes de escribir nada.

La lectura no depende del formato exacto del archivo: las cabeceras se buscan
por su texto en lugar de fijar numeros de fila.
"""

from __future__ import annotations

import sqlite3
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from mukuwareru.nucleo.modelos.entidades import OrigenSesion, TipoSesion
from mukuwareru.nucleo.repositorios import (
    RepositorioMaterias,
    RepositorioModulos,
    RepositorioSesiones,
)
from mukuwareru.nucleo.repositorios.base import transaccion
from mukuwareru.utilidades.registro import obtener

_log = obtener(__name__)

# Cabeceras que identifican cada hoja, en minusculas.
_CAB_TEMA = ("tema", "materia")
_CAB_MODULO = ("modulo", "módulo", "reading", "modulo / reading", "módulo / reading")
_CAB_COMPLETADO = ("completado", "completada", "hecho")
_CAB_FECHA = ("fecha", "dia", "día")
_CAB_HORAS = ("horas", "tiempo")

# Filas de totales que las hojas de calculo arrastran y que no son datos.
_PREFIJOS_TOTAL = ("total", "suma", "resumen")


@dataclass(slots=True)
class ModuloImportado:
    """Una fila valida de la hoja de modulos."""

    materia: str
    nombre: str
    completado: bool


@dataclass(slots=True)
class SesionImportada:
    """Una fila valida de la hoja de registro diario."""

    fecha: datetime
    segundos: int


@dataclass(slots=True)
class InformeImportacion:
    """Vista previa de lo que se va a importar. No escribe nada."""

    modulos: list[ModuloImportado] = field(default_factory=list)
    sesiones: list[SesionImportada] = field(default_factory=list)
    descartadas: list[str] = field(default_factory=list)
    hojas_ausentes: list[str] = field(default_factory=list)

    @property
    def materias(self) -> list[str]:
        """Materias distintas, en el orden en que aparecen en el archivo."""
        vistas: dict[str, None] = {}
        for modulo in self.modulos:
            vistas.setdefault(modulo.materia, None)
        return list(vistas)

    @property
    def completados(self) -> int:
        """Cuantos modulos vienen marcados como completados."""
        return sum(1 for m in self.modulos if m.completado)

    @property
    def horas(self) -> float:
        """Horas totales del registro diario."""
        return sum(s.segundos for s in self.sesiones) / 3600


@dataclass(frozen=True, slots=True)
class ResultadoImportacion:
    """Que se escribio realmente en la base de datos."""

    materias_creadas: int
    modulos_creados: int
    modulos_actualizados: int
    sesiones_creadas: int
    sesiones_omitidas: int = 0


class ServicioImportacion:
    """Lee un Excel de progreso y lo vuelca en un proyecto."""

    def __init__(self, conexion: sqlite3.Connection) -> None:
        self._cx = conexion
        self._materias = RepositorioMaterias(conexion)
        self._modulos = RepositorioModulos(conexion)
        self._sesiones = RepositorioSesiones(conexion)

    # -- Paso 1: analizar ---------------------------------------------------

    def analizar(self, ruta: Path) -> InformeImportacion:
        """Lee el archivo y devuelve lo que se importaria, sin escribir nada."""
        informe = InformeImportacion()
        libro = load_workbook(ruta, data_only=True, read_only=True)
        try:
            hoja_modulos = _buscar_hoja(libro, _CAB_TEMA, _CAB_MODULO)
            if hoja_modulos is None:
                informe.hojas_ausentes.append("modulos")
            else:
                _leer_modulos(hoja_modulos, informe)

            hoja_registro = _buscar_hoja(libro, _CAB_FECHA, _CAB_HORAS)
            if hoja_registro is None:
                informe.hojas_ausentes.append("registro diario")
            else:
                _leer_sesiones(hoja_registro, informe)
        finally:
            libro.close()

        _log.info(
            "Analisis de %s: %d modulos (%d completados), %d sesiones, %d filas descartadas",
            ruta.name,
            len(informe.modulos),
            informe.completados,
            len(informe.sesiones),
            len(informe.descartadas),
        )
        return informe

    # -- Paso 2: aplicar ----------------------------------------------------

    def aplicar(self, proyecto_id: int, informe: InformeImportacion) -> ResultadoImportacion:
        """Vuelca el informe en el proyecto, dentro de una unica transaccion.

        Es idempotente: los modulos se identifican por
        ``(proyecto, materia, nombre)`` y las sesiones por su fecha, de modo que
        reimportar el mismo archivo actualiza el estado en vez de duplicar filas.
        """
        materias_creadas = modulos_creados = modulos_actualizados = 0
        sesiones_creadas = sesiones_omitidas = 0

        with transaccion(self._cx):
            existentes = {m.nombre: m for m in self._materias.listar(proyecto_id)}
            indice_modulos: dict[tuple[int, str], tuple[int, bool]] = {
                (materia_id, modulo.nombre): (modulo.id, modulo.completado)
                for materia_id, lista in self._modulos.listar_del_proyecto(proyecto_id).items()
                for modulo in lista
            }

            orden_materia = len(existentes)
            orden_modulo: dict[int, int] = {}

            for fila in informe.modulos:
                materia = existentes.get(fila.materia)
                if materia is None:
                    materia = self._materias.crear(
                        proyecto_id, fila.materia, orden=orden_materia
                    )
                    existentes[fila.materia] = materia
                    orden_materia += 1
                    materias_creadas += 1

                clave = (materia.id, fila.nombre)
                if (previo := indice_modulos.get(clave)) is None:
                    siguiente = orden_modulo.get(materia.id, 0)
                    self._modulos.crear(
                        materia.id, fila.nombre, orden=siguiente, completado=fila.completado
                    )
                    orden_modulo[materia.id] = siguiente + 1
                    modulos_creados += 1
                elif previo[1] != fila.completado:
                    self._modulos.marcar(previo[0], fila.completado)
                    modulos_actualizados += 1

            ya_importadas = self._sesiones.fechas_importadas(proyecto_id)
            for sesion in informe.sesiones:
                if (fecha := sesion.fecha.date().isoformat()) in ya_importadas:
                    sesiones_omitidas += 1
                    continue
                self._sesiones.crear(
                    proyecto_id,
                    tipo=TipoSesion.TRABAJO,
                    origen=OrigenSesion.IMPORTADA,
                    inicio=sesion.fecha,
                    duracion_seg=sesion.segundos,
                    completada=True,
                )
                ya_importadas.add(fecha)
                sesiones_creadas += 1

        resultado = ResultadoImportacion(
            materias_creadas=materias_creadas,
            modulos_creados=modulos_creados,
            modulos_actualizados=modulos_actualizados,
            sesiones_creadas=sesiones_creadas,
            sesiones_omitidas=sesiones_omitidas,
        )
        _log.info("Importacion aplicada: %s", resultado)
        return resultado


# ---------------------------------------------------------------------------
# Lectura de hojas
# ---------------------------------------------------------------------------


def _normalizar(valor: Any) -> str:
    return str(valor).strip().lower() if valor is not None else ""


def _indice_cabecera(fila: tuple[Any, ...], alternativas: tuple[str, ...]) -> int | None:
    """Posicion de la primera celda que coincide con alguna alternativa."""
    for indice, celda in enumerate(fila):
        texto = _normalizar(celda)
        if texto and any(texto == a or texto.startswith(a) for a in alternativas):
            return indice
    return None


def _buscar_hoja(libro: Any, *cabeceras: tuple[str, ...]) -> Any:
    """Primera hoja que contiene una fila con todas las cabeceras pedidas."""
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows(min_row=1, max_row=15, values_only=True):
            if all(_indice_cabecera(fila, c) is not None for c in cabeceras):
                return hoja
    return None


def _fila_cabecera(hoja: Any, *cabeceras: tuple[str, ...]) -> tuple[int, list[int]]:
    """Numero de fila de la cabecera y posicion de cada columna pedida."""
    for numero, fila in enumerate(
        hoja.iter_rows(min_row=1, max_row=15, values_only=True), start=1
    ):
        indices = [_indice_cabecera(fila, c) for c in cabeceras]
        if all(i is not None for i in indices):
            return numero, [i for i in indices if i is not None]
    return 0, []


def _es_total(texto: str) -> bool:
    """Detecta las filas de totales que arrastran las hojas de calculo."""
    return any(texto.lower().startswith(p) for p in _PREFIJOS_TOTAL)


def _celda(fila: tuple[Any, ...], indice: int) -> Any:
    return fila[indice] if indice < len(fila) else None


def _leer_modulos(hoja: Any, informe: InformeImportacion) -> None:
    """Vuelca la hoja de modulos en el informe."""
    cabecera, (col_tema, col_modulo) = _fila_cabecera(hoja, _CAB_TEMA, _CAB_MODULO)
    if not cabecera:
        return
    _, columnas_completado = _fila_cabecera(hoja, _CAB_COMPLETADO)
    col_completado = columnas_completado[0] if columnas_completado else None

    for fila in hoja.iter_rows(min_row=cabecera + 1, values_only=True):
        tema = str(_celda(fila, col_tema) or "").strip()
        nombre = str(_celda(fila, col_modulo) or "").strip()
        if not tema and not nombre:
            continue
        if _es_total(tema) or _es_total(nombre):
            informe.descartadas.append(f"Fila de totales: {tema or nombre}")
            continue
        if not tema or not nombre:
            informe.descartadas.append(f"Fila incompleta: {tema or nombre}")
            continue

        bruto = _celda(fila, col_completado) if col_completado is not None else None
        informe.modulos.append(
            ModuloImportado(materia=tema, nombre=nombre, completado=_a_booleano(bruto))
        )


def _leer_sesiones(hoja: Any, informe: InformeImportacion) -> None:
    """Vuelca la hoja de registro diario en el informe."""
    cabecera, (col_fecha, col_horas) = _fila_cabecera(hoja, _CAB_FECHA, _CAB_HORAS)
    if not cabecera:
        return

    for fila in hoja.iter_rows(min_row=cabecera + 1, values_only=True):
        bruto_fecha = _celda(fila, col_fecha)
        bruto_horas = _celda(fila, col_horas)
        if bruto_fecha is None and bruto_horas is None:
            continue
        if isinstance(bruto_fecha, str) and _es_total(bruto_fecha):
            informe.descartadas.append(f"Fila de totales: {bruto_fecha}")
            continue
        if not isinstance(bruto_fecha, datetime):
            if bruto_fecha is not None:
                informe.descartadas.append(f"Fecha no reconocida: {bruto_fecha!r}")
            continue
        if not isinstance(bruto_horas, int | float) or bruto_horas <= 0:
            continue

        informe.sesiones.append(
            SesionImportada(fecha=bruto_fecha, segundos=round(float(bruto_horas) * 3600))
        )


def _a_booleano(valor: Any) -> bool:
    """Interpreta las muchas formas de escribir «si» en una hoja de calculo."""
    if isinstance(valor, bool):
        return valor
    if isinstance(valor, int | float):
        return bool(valor)
    return _normalizar(valor) in {"si", "sí", "true", "verdadero", "x", "1", "yes", "ok"}
