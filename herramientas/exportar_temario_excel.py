"""Reescribe la hoja «Modulos» de CFA.xlsx con el temario de la aplicacion.

    python herramientas/exportar_temario_excel.py            # simulacion
    python herramientas/exportar_temario_excel.py --aplicar  # escribe

La direccion es **de la aplicacion al Excel**, y no al reves. Despues de la
importacion inicial la fuente oficial es la base de datos (§9 de PROJECT.md); el
Excel se queda como copia de consulta. Sin este paso, reimportar el archivo
viejo resucitaria los modulos que se acaban de corregir, porque el importador
solo anade.

Dos cuidados con el archivo:

* Se escribe **el booleano de verdad**, no la formula ``=FALSE()`` que habia.
  openpyxl no calcula formulas, asi que una celda con formula recien escrita no
  tiene valor cacheado y ``load_workbook(data_only=True)`` —lo que usa el
  importador— la leeria como vacia: todos los completados se perderian.
* El Dashboard usa ``COUNTIF`` sobre columnas enteras, asi que crecer de 75 a 93
  filas no le afecta. Lo unico que hay que recolocar es la fila de totales.
"""

from __future__ import annotations

import argparse
import shutil
import sys
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from herramientas._comun import base_instalada  # noqa: E402
from mukuwareru.nucleo.bd import conexion as bd  # noqa: E402
from mukuwareru.nucleo.repositorios import (  # noqa: E402
    RepositorioMaterias,
    RepositorioModulos,
    RepositorioProyectos,
)

PROYECTO = "CFA Level I"
HOJA = "Módulos"
FILA_CABECERA = 5
COL_TEMA, COL_MODULO, COL_HECHO = 2, 3, 4


def leer_temario(ruta_base: Path) -> list[tuple[str, str, bool]]:
    """Filas ``(tema, modulo, completado)`` del proyecto, en su orden."""
    conexion = bd.abrir(ruta_base)
    try:
        proyecto = next(
            (p for p in RepositorioProyectos(conexion).listar() if p.nombre == PROYECTO),
            None,
        )
        if proyecto is None:
            raise SystemExit(f"No existe el proyecto «{PROYECTO}».")

        materias = RepositorioMaterias(conexion)
        modulos = RepositorioModulos(conexion)
        return [
            (materia.nombre, modulo.nombre, modulo.completado)
            for materia in materias.listar(proyecto.id)
            for modulo in modulos.listar(materia.id)
        ]
    finally:
        conexion.close()


def escribir(ruta_excel: Path, filas: list[tuple[str, str, bool]]) -> tuple[int, int]:
    """Vuelca las filas en la hoja. Devuelve ``(primera, ultima)`` fila escrita."""
    libro = load_workbook(ruta_excel)
    hoja = libro[HOJA]

    # La fila de totales se busca por contenido: fijar el numero de fila
    # romperia en cuanto el temario cambie de tamano otra vez.
    fila_total = next(
        (
            fila
            for fila in range(FILA_CABECERA, hoja.max_row + 1)
            if str(hoja.cell(fila, COL_TEMA).value or "").lower().startswith("total")
        ),
        None,
    )

    primera = FILA_CABECERA + 1
    plantilla = [
        copy(hoja.cell(primera, columna)._style)
        for columna in (COL_TEMA, COL_MODULO, COL_HECHO)
    ]

    # La fila de totales se lee y se vacia **antes** de escribir los datos: con
    # 93 modulos donde habia 75, las filas nuevas la pisan, y capturarla despues
    # guardaria un nombre de tema como si fuera la etiqueta del total.
    total_texto = ""
    total_estilos: list[object] = []
    if fila_total is not None:
        etiqueta = hoja.cell(fila_total, COL_TEMA)
        formula = hoja.cell(fila_total, COL_HECHO)
        total_texto = str(etiqueta.value or "")
        total_estilos = [copy(etiqueta._style), copy(formula._style)]
        etiqueta.value = formula.value = None

    for indice, (tema, modulo, hecho) in enumerate(filas):
        fila = primera + indice
        for columna, valor, estilo in zip(
            (COL_TEMA, COL_MODULO, COL_HECHO), (tema, modulo, hecho), plantilla, strict=True
        ):
            celda = hoja.cell(fila, columna)
            celda.value = valor
            celda._style = copy(estilo)

    ultima = primera + len(filas) - 1

    # Limpia lo que sobre del temario anterior.
    for fila in range(ultima + 1, hoja.max_row + 1):
        for columna in (COL_TEMA, COL_MODULO, COL_HECHO):
            hoja.cell(fila, columna).value = None

    if total_texto:
        destino = ultima + 2
        nueva_etiqueta = hoja.cell(destino, COL_TEMA)
        nueva_etiqueta.value = total_texto
        nueva_etiqueta._style = total_estilos[0]
        nueva_formula = hoja.cell(destino, COL_HECHO)
        nueva_formula.value = (
            f'=COUNTIF(D{primera}:D{ultima},TRUE())&" / "&COUNTA(D{primera}:D{ultima})'
        )
        nueva_formula._style = total_estilos[1]

    libro.save(ruta_excel)
    return primera, ultima


def main() -> int:
    """Punto de entrada."""
    analizador = argparse.ArgumentParser(description=__doc__)
    analizador.add_argument("--aplicar", action="store_true", help="escribe el archivo")
    analizador.add_argument("--excel", type=Path, default=RAIZ / "CFA.xlsx")
    analizador.add_argument(
        "--base", type=Path, default=None, help="base de datos (por defecto, la instalada)"
    )
    argumentos = analizador.parse_args()

    if not argumentos.excel.exists():
        print(f"No se encuentra {argumentos.excel}")
        return 1

    filas = leer_temario(argumentos.base or base_instalada())
    hechos = sum(1 for _t, _m, h in filas if h)
    print(f"Temario en la aplicacion: {len(filas)} modulos, {hechos} completados")

    por_tema: dict[str, int] = {}
    for tema, _modulo, _hecho in filas:
        por_tema[tema] = por_tema.get(tema, 0) + 1
    for tema, cuantos in por_tema.items():
        print(f"  {tema:<26} {cuantos:>2}")

    if not argumentos.aplicar:
        print("\nSimulacion. Repite con --aplicar para escribir el Excel.")
        return 0

    sello = datetime.now().strftime("%Y%m%d-%H%M%S")
    respaldo = argumentos.excel.with_suffix(f".respaldo-{sello}.xlsx")
    shutil.copy(argumentos.excel, respaldo)
    print(f"\nRespaldo: {respaldo.name}")

    primera, ultima = escribir(argumentos.excel, filas)
    print(f"Escritas las filas {primera}-{ultima} de «{HOJA}» en {argumentos.excel.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
